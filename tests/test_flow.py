"""흐름 테스트 · 가짜 API로 CSV → 조회 → 판정 → 결과 파일까지 끝까지 돌린다. 인터넷 불필요

원본 JS 대조(test_js_parity)만 Node.js가 필요하고, 없으면 건너뛴다.
"""
import csv
import glob
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
sys.path.insert(0, HERE)

import fakeapi  # noqa: E402
import sample  # noqa: E402
from test_units import KEYS, TmpCase  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pipeline import cli, service  # noqa: E402
from pipeline.errors import PipelineError, StageStopped, UsageError  # noqa: E402
from pipeline.selection import FilterSet  # noqa: E402
from pipeline.compat import enc_uri  # noqa: E402

QUIET = lambda *a, **k: None  # noqa: E731


class FlowTest(TmpCase):
    def setUp(self):
        super().setUp()
        self.src = sample.write(os.path.join(self.tmp, '콜리스트_미추홀구.csv'))
        self.work, self.out = os.path.join(self.tmp, 'work'), os.path.join(self.tmp, 'results')

    def core(self, t=None, **kw):
        kw.setdefault('filter_sets', [FilterSet(gu=['미추홀구'])])
        return service.run_core([self.src], work=self.work, out=self.out, transport=t or fakeapi.FakeTransport(), date='2026-10-01', log=QUIET, **kw)

    def test_full_run_outputs(self):
        r = self.core()
        d = r['dir']
        g = os.path.join(d, '미추홀구_관교동')
        files = sorted(os.path.basename(p) for p in glob.glob(g + '/*'))
        self.assertIn('콜리스트_미추홀구_관교동_재검색결과_2026-10-01.xlsx', files)
        self.assertIn('콜리스트_미추홀구_관교동_재검색결과_2026-10-01.csv', files)
        self.assertEqual(sum(1 for f in files if f.endswith('.html')), 5)
        with open(os.path.join(g, '콜리스트_미추홀구_관교동_재검색결과_2026-10-01.csv'), encoding='utf-8-sig') as f:
            rows = list(csv.DictReader(f))
        self.assertEqual(len(rows[0]), 77)                       # 원본 28열 + 새 열 49개
        self.assertEqual({r['buildingRosterIdx'] for r in rows}, {'101', '102', '103', '104', '105'})   # 106은 조회 대상 아님
        by = {r['buildingRosterIdx']: r for r in rows}
        self.assertEqual(by['101']['건물명_확정'], '동부아파트')
        self.assertEqual(by['104']['건물명_상태'], '미해결')
        self.assertTrue(by['105']['대리1_번호'])
        self.assertTrue(by['101']['건물명_확정_링크'].startswith('https://map.kakao.com/'))
        with open(os.path.join(d, 'summary.json'), encoding='utf-8') as f:
            s = json.load(f)
        self.assertEqual([x['동'] for x in s['groups']], ['관교동', '학익동'])
        self.assertEqual(s['settings']['extra_words'], {'주안동': ['관교']})
        with open(glob.glob(g + '/*사람확인필요*.html')[0], encoding='utf-8') as f:
            html = f.read()
        self.assertIn('<title>미추홀구 관교동 사람확인필요</title>', html)

    def test_formula_guard(self):
        r = self.core()
        x = glob.glob(os.path.join(r['dir'], '미추홀구_관교동', '*.xlsx'))[0]
        wb = load_workbook(x)
        for ws in wb.worksheets:
            f = [c for row in ws.iter_rows() for c in row if c.data_type == 'f']
            if ws.title == '지표':
                self.assertTrue(f)
            else:
                self.assertEqual(f, [], ws.title)
        texts = [c.value for row in wb['근거_주소일치장소'].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertEqual(len(texts), 1)

    def test_resume_makes_no_calls_and_new_folder(self):
        r1 = self.core()
        t2 = fakeapi.FakeTransport()
        r2 = self.core(t2)
        self.assertEqual(r1['job_id'], r2['job_id'])
        self.assertEqual(t2.count, {})
        self.assertNotEqual(r1['dir'], r2['dir'])
        self.assertTrue(os.path.exists(r1['dir']))

    def test_job_id_changes(self):
        base = service.Plan([self.src], [FilterSet(gu=['미추홀구'])], False, '').job_id
        self.assertEqual(len(base), 24)
        self.assertNotEqual(base, service.Plan([self.src], [FilterSet(gu=['미추홀구'])], False, 'again').job_id)
        self.assertNotEqual(base, service.Plan([self.src], [FilterSet(gu=['미추홀구'])], True, '').job_id)
        self.assertNotEqual(base, service.Plan([self.src], [FilterSet(gu=['미추홀구'], rank=['B'])], False, '').job_id)
        same = service.Plan([self.src], [FilterSet(gu=['미추홀구', '미추홀구'])], False, '').job_id
        self.assertEqual(base, same)

    def test_stop_then_resume_and_114_gate(self):
        bad = {fakeapi.kaddr('인천광역시 미추홀구 ' + a): {'status': 400, 'body': ''} for a in
               ('주승로 247', '인하로411번길 49', '관교동 500', '경원대로 627')}
        with self.assertRaises(StageStopped):
            self.core(fakeapi.FakeTransport(overrides=bad), workers=1)
        jid = service.Plan([self.src], [FilterSet(gu=['미추홀구'])], False, '').job_id
        with self.assertRaises(UsageError):
            service.run_114(jid, work=self.work, out=self.out, transport=fakeapi.FakeTransport(), log=QUIET)
        t = fakeapi.FakeTransport()
        r = self.core(t)
        self.assertGreater(t.count.get('kakao', 0), 0)
        r114 = service.run_114(r['job_id'][:8], 'no_rep', work=self.work, out=self.out, transport=fakeapi.FakeTransport(), date='2026-10-01', log=QUIET)
        self.assertIn(os.sep + '114' + os.sep, r114['dir'])
        m = service.status(r['job_id'], self.work, log=QUIET)
        self.assertTrue(m['c114']['complete'])

    def test_failed_geocode_is_retried_next_run(self):
        """기존 merge.py 는 NG.json 에 키가 있으면 대기열에서 뺀다. 실패 행을 넘기지 않아야 다음 실행에서 다시 조회된다"""
        geo = '/map-geocode/v2/geocode?query=' + enc_uri('인천광역시 미추홀구 경원대로640번길 6-39')
        r1 = self.core(fakeapi.FakeTransport(overrides={geo: {'status': 500, 'body': ''}}))
        self.assertGreater(r1['summary']['errors_remaining'], 0)
        t2 = fakeapi.FakeTransport()
        r2 = self.core(t2)
        self.assertIn(geo, t2.calls)
        self.assertNotIn('kakao', t2.count)
        self.assertEqual(r2['summary']['errors_remaining'], 0)

    def test_run_py_reads_env_file(self):
        """F5 「기본 조회」(run.py)도 .env 의 키를 읽어야 한다 (전에는 run_config.py 만 읽었음)"""
        from pipeline import credentials
        env = os.path.join(self.tmp, '.env')
        with open(env, 'w', encoding='utf-8') as f:
            f.write('# 주석\nKAKAO_REST_KEY=(from-env-file-123)\n')
        old_default, old_key = credentials.DEFAULT_ENV, os.environ.pop('KAKAO_REST_KEY', None)
        credentials.DEFAULT_ENV = env
        try:
            with redirect_stdout(io.StringIO()), self.assertRaises(SystemExit) as cm:
                cli.main(['status', '--work', self.work])
            self.assertEqual(cm.exception.code, 0)
            self.assertEqual(os.environ.get('KAKAO_REST_KEY'), 'from-env-file-123')
        finally:
            credentials.DEFAULT_ENV = old_default
            os.environ.pop('KAKAO_REST_KEY', None)
            if old_key is not None:
                os.environ['KAKAO_REST_KEY'] = old_key

    def test_missing_key_only_when_needed(self):
        os.environ.pop('NCP_HUB_ID')
        with self.assertRaises(PipelineError) as c:
            self.core()
        self.assertIn('NCP_HUB_ID', str(c.exception))
        os.environ['NCP_HUB_ID'] = KEYS['NCP_HUB_ID']
        self.core()
        for k in KEYS:
            os.environ.pop(k)
        service.export(service.Plan([self.src], [FilterSet(gu=['미추홀구'])], False, '').job_id, work=self.work, out=self.out, log=QUIET)

    def test_leaked_key_is_caught(self):
        os.environ['KAKAO_REST_KEY'] = '동부아파트 관리사무소'   # 결과에 들어가는 글자를 키로 가정
        with self.assertRaises(PipelineError) as c:
            self.core()
        self.assertIn('API 키', str(c.exception))

    def test_cli_exit_codes(self):
        e = io.StringIO()
        with redirect_stdout(io.StringIO()), redirect_stderr(e):
            code = cli.guarded(lambda: cli.dispatch(cli.build_parser().parse_args(['list', '--src', self.src, '--gu', '미추홀'])))
        self.assertEqual(code, 2)
        self.assertIn('미추홀구', e.getvalue())
        with redirect_stdout(io.StringIO()):
            self.assertEqual(cli.guarded(lambda: cli.dispatch(cli.build_parser().parse_args(['list', '--src', self.src, '--gu', '미추홀구,중구'][:5]))), 2)
            self.assertEqual(cli.guarded(lambda: cli.dispatch(cli.build_parser().parse_args(['list', '--src', self.src]))), 0)

        def cancel():
            raise KeyboardInterrupt
        with redirect_stderr(io.StringIO()):
            self.assertEqual(cli.guarded(cancel), 130)
        r = subprocess.run([sys.executable, os.path.join(ROOT, 'run.py'), 'run', '--src', self.src, '--workers', '9'], capture_output=True, text=True)
        self.assertEqual(r.returncode, 2)
        r = subprocess.run([sys.executable, os.path.join(ROOT, 'run.py'), 'nope'], capture_output=True, text=True)
        self.assertEqual(r.returncode, 2)


class RunConfigTest(TmpCase):
    def test_config_union_env_and_last_job(self):
        import run_config
        sample.write(os.path.join(self.tmp, 'a.csv'), sample.ROWS + sample.ROWS_JUNG)
        cfg = os.path.join(self.tmp, '조회설정.toml')
        with open(cfg, 'w', encoding='utf-8') as f:
            f.write('src = ["a.csv"]\n[["조회"]]\n"구" = ["중구"]\n[[query]]\ngu = ["미추홀구"]\nrank = ["b"]\n')
        c = run_config.load_config(cfg)
        self.assertEqual(len(c['filters']), 2)
        self.assertEqual(c['filters'][1].rank, ['B'])
        plan = service.Plan(c['src'], c['filters'], False, '')
        self.assertEqual(sorted(plan.df['buildingRosterIdx']), ['101', '103', '301'])
        with open(cfg, 'a', encoding='utf-8') as f:
            f.write('"오타" = 1\n')
        with self.assertRaises(UsageError):
            run_config.load_config(cfg)
        bad = os.path.join(self.tmp, 'bad.toml')
        with open(bad, 'w', encoding='utf-8') as f:
            f.write('src = ["a.csv"]\n[[조회]]\n구 = ["중구"]\n')
        with self.assertRaises(UsageError) as c:
            run_config.load_config(bad)
        self.assertIn('따옴표', str(c.exception))
        shipped = run_config.load_config(os.path.join(ROOT, '조회설정.toml'))
        self.assertEqual([f.gu for f in shipped['filters']], [['중구'], ['미추홀구']])
        env = os.path.join(self.tmp, '.env')
        with open(env, 'w', encoding='utf-8') as f:
            f.write('# 주석\nNCP_GEO_ID="from-env-file"\nEMPTY=\n')
        os.environ.pop('NCP_GEO_ID')
        self.assertEqual(run_config.load_env(env), ['NCP_GEO_ID'])
        self.assertEqual(os.environ['NCP_GEO_ID'], 'from-env-file')
        old = run_config.LAST
        run_config.LAST = os.path.join(self.tmp, '.last_job.json')
        try:
            run_config.remember('abc123', cfg)
            self.assertEqual(run_config.last_job(), 'abc123')
        finally:
            run_config.LAST = old


@unittest.skipUnless(shutil.which('node'), 'Node.js 없음')
class JsParityTest(unittest.TestCase):
    """원본 collector.js·naver.js 와 파이썬 이식본을 같은 가짜 응답으로 돌려 결과와 요청 주소를 비교"""

    def test_same_as_original_js(self):
        from pipeline.providers.kakao import Kakao, proc_a
        from pipeline.providers.kb import KB
        from pipeline.providers.naver import Naver
        from pipeline.providers.one114 import One114, proc_c
        from pipeline.providers.zippoom import Zippoom
        from pipeline.secondary import proc_b
        routes = fakeapi.routes()
        inp = {'routes': routes, 'defaults': fakeapi.DEFAULTS, 'rows': fakeapi.ROWS_A, 'geo': fakeapi.GEO_ROWS, 'verify': fakeapi.VERIFY_ROWS,
               'q114': fakeapi.Q114_ROWS, 'gu': '미추홀구', 'dong': '관교동'}
        with tempfile.NamedTemporaryFile('w', suffix='.json', delete=False, encoding='utf-8') as f:
            json.dump(inp, f, ensure_ascii=False)
        try:
            r = subprocess.run(['node', os.path.join(HERE, 'reference', 'harness.js'), f.name], capture_output=True, text=True, encoding='utf-8')
        finally:
            os.unlink(f.name)
        self.assertEqual(r.returncode, 0, r.stderr)
        js = json.loads(r.stdout)
        t = fakeapi.FakeTransport(routes)
        k = Kakao(t, 'KEY')
        py = {'A': {}, 'B': {}, 'C': {}, 'NG': {}, 'NV': {}}
        for b in fakeapi.ROWS_A:
            py['A'][str(b[0])] = proc_a(k, '미추홀구', b)
        for b in fakeapi.ROWS_A:
            py['B'][str(b[0])] = proc_b(k, KB(t), Zippoom(t), t.sleep, '미추홀구', b, py['A'][str(b[0])])
        one = One114(t)
        for q in fakeapi.Q114_ROWS:
            py['C'][str(q[0])] = proc_c(one, '미추홀구', q, lambda: 0)
        nv = Naver(t, '미추홀구', geo_key=('ID', 'SECRET'), hub_key=('ID', 'SECRET'))
        for g in fakeapi.GEO_ROWS:
            py['NG'][str(g[0])] = nv.proc_geo(g)
        for v in fakeapi.VERIFY_ROWS:
            py['NV'][v[0]] = nv.proc_verify(v)
        for side in (js, py):
            for c in side['C'].values():
                c.pop('t', None)
        norm = lambda o: json.loads(json.dumps(o, ensure_ascii=False))  # noqa: E731
        for st in ('A', 'B', 'C', 'NG', 'NV'):
            self.assertEqual(norm(js[st]), norm(py[st]), st)
        self.assertEqual(sorted(js['calls']), sorted(t.calls))


if __name__ == '__main__':
    unittest.main()
