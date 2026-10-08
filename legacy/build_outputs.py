"""결과 파일 생성 · 결과 엑셀 1개와 HTML 목록 5개를 만든다

사용 예
  python build_outputs.py --src 콜리스트.csv --work work --gu 미추홀구 --dong 주안동 --out out --date 2026-09-28

만드는 파일 · --out 폴더
  콜리스트_{구}_{동}_재검색결과_{날짜}.xlsx
      지표 · 결과 · 근거_주소일치장소 · 근거_KB부동산·집품 · 검수_후보번호 시트
  {구}_{동}_사람확인필요_목록_{날짜}.html
  {구}_{동}_건물명번호_둘다미해결_목록_{날짜}.html
  {구}_{동}_건물명만미해결_목록_{날짜}.html
  {구}_{동}_번호만미해결_목록_{날짜}.html
  {구}_{동}_채운정보_전체목록_{날짜}.html

HTML 정렬 규칙 · 모든 목록 공통
  맨 위에 사람 확인 필요 구역, 그 아래 나머지. 두 구역 모두 등급 S A B C D 순, 등급 없음 맨 끝, 같은 등급은 idx 순
  구역 제목과 등급 소제목에 건수 표시. 한 행은 한 구역에만 들어감

원본 파일은 읽기만 하고, 결과는 항상 새 파일로 씀
"""
import argparse, html, json, os, re, subprocess, shutil
from urllib.parse import quote
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from common import load, s, sort_key, GRADE_ORDER, short_addr

F = 'Arial'

def links_for(x):
    q = x['road'] or x['jibun']
    g = x.get('g')
    zq = q.replace('인천광역시', '인천')
    kb = x.get('kb')
    return {
        '카카오맵': 'https://map.kakao.com/?q=' + quote(q),
        '카카오 로드뷰': f'https://map.kakao.com/link/roadview/{g[1]},{g[0]}' if g else '',
        '네이버지도': 'https://map.naver.com/p/search/' + quote(q),
        'KB부동산': 'https://kbland.kr/c/' + str(kb['no']) if kb and kb.get('no') else '',
        '집품': 'https://zippoom.com/search/' + quote(zq) + '?searchWithAddress=false',
        '114On': 'https://www.114.co.kr/search/result/all?query=' + quote(q),
    }

def proxy_rows(x):
    out = []
    for p in x.get('st') or []:
        out.append(('건물 내 점포 · ' + p['role'], p['n'], p['p'], p['src'] + ' · ' + p['c'], p['link']))
    for p in x.get('re') or []:
        out.append((f"주변 부동산 · {p['d']}m", p['n'], p['p'], '카카오맵 · 반경 300m 부동산', p['link']))
    return out

def status(x):
    if not x['nt']: ns = '대상아님(기존 건물명 있음)'
    elif x.get('nm'): ns = '채움 · ' + x['nm']['st']
    else: ns = '미해결'
    if not x['pt']: ps = '대상아님(기존 번호 있음)'
    elif x.get('ph'): ps = '채움 · ' + x['ph']['src']
    elif x.get('st') or x.get('re'): ps = '대리 연락처만 확보'
    else: ps = '미해결'
    return ns, ps

# ---------------- 엑셀 ----------------
def build_xlsx(df, R, gu, dong, out, date):
    orig = list(df.columns)
    new = ['건물명_확정', '건물명_상태', '건물명_출처', '대표번호_확정', '번호_상태', '번호_출처']
    for k in range(1, 6): new += [f'대리{k}_구분', f'대리{k}_건물명', f'대리{k}_번호', f'대리{k}_출처']
    new += ['검수_후보번호', '주소일치_카카오장소수', '주소일치_114On점포수', '검색기록', '근거_출처', '근거_링크',
            '사람확인필요', '사람확인_사유', '시설건물', '시설_종류', '시설_판단근거']
    rows = []
    for _, row in df.iterrows():
        v = {c: row[c] for c in orig}; L = {}
        x = R.get(int(row['buildingRosterIdx']))
        if x is None:
            v.update({'건물명_확정': row['건물명'], '건물명_상태': '대상아님(기존 건물명 있음)', '대표번호_확정': row['대표번호'], '번호_상태': '대상아님(기존 번호 있음)'})
            rows.append((v, L)); continue
        ns, ps = status(x); ev_s, ev_l = [], []
        v['건물명_상태'] = ns
        if x['nt'] and x.get('nm'):
            v['건물명_확정'] = x['nm']['v']; v['건물명_출처'] = x['nm']['src']; L['건물명_확정'] = x['nm']['link']
            ev_s.append('[건물명] ' + x['nm']['src']); ev_l.append(x['nm']['link'])
        else:
            v['건물명_확정'] = row['건물명']
        v['번호_상태'] = ps
        if x['pt'] and x.get('ph'):
            p = x['ph']; v['대표번호_확정'] = p['v']; v['번호_출처'] = p['name']; L['대표번호_확정'] = p['link']
            ev_s.append('[대표번호] ' + p['src'] + ' ' + p['name']); ev_l.append(p['link'])
        else:
            v['대표번호_확정'] = row['대표번호'] if not x['pt'] else ''
        pr = proxy_rows(x)
        for k in range(1, 6):
            if k <= len(pr):
                g, n, p, src, link = pr[k - 1]
                v[f'대리{k}_구분'] = g; v[f'대리{k}_건물명'] = n; v[f'대리{k}_번호'] = p; v[f'대리{k}_출처'] = src
                L[f'대리{k}_건물명'] = link; L[f'대리{k}_번호'] = link
        v['검수_후보번호'] = ' / '.join(f"{h['v']} {h['name']} · {h['reason']}" for h in x.get('held') or [])
        v['주소일치_카카오장소수'] = x.get('km', 0); v['주소일치_114On점포수'] = x.get('tm', 0)
        v['검색기록'] = ' | '.join(l for l in x.get('log') or [] if l)
        if not ev_s:
            ev_s.append('카카오 주소검색·장소검색 결과, 건물 자체 정보 없음'); ev_l.append('https://map.kakao.com/?q=' + quote(x['road'] or x['jibun']))
        v['근거_출처'] = ' / '.join(ev_s); v['근거_링크'] = ev_l[0]; L['근거_링크'] = ev_l[0]
        v['사람확인필요'] = '예' if x.get('human') else ''
        v['사람확인_사유'] = ' · '.join(x.get('flags') or [])
        lab, why, fl = (x.get('fac') or ['', '', ''])
        v['시설건물'] = '예' if lab else ''; v['시설_종류'] = lab; v['시설_판단근거'] = why
        if fl: L['시설_판단근거'] = fl
        rows.append((v, L))
    cols = orig + new
    wb = Workbook(); ws = wb.active; sheet = f'{dong}_결과'[:31]; ws.title = sheet
    fo, fn, fp = PatternFill('solid', fgColor='D9D9D9'), PatternFill('solid', fgColor='DCE6F1'), PatternFill('solid', fgColor='FFF2CC')
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c); cell.font = Font(name=F, bold=True)
        cell.fill = fp if c.startswith('대리') else (fn if c in new else fo)
    lf = Font(name=F, color='0563C1', underline='single')
    for i, (v, L) in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            val = v.get(c, '')
            if isinstance(val, float) and pd.isna(val): val = ''
            cell = ws.cell(i, j, val); cell.font = Font(name=F)
            if L.get(c) and val != '': cell.hyperlink = L[c]; cell.font = lf
    ws.freeze_panes = 'J2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows) + 1}'
    for j, c in enumerate(cols, 1):
        w = 12
        if c in ('건물명', '건물명_확정', '도로명주소', '지번주소') or c.endswith('_건물명'): w = 26
        if c in ('검색기록', '근거_출처', '연락처 전체', '검수_후보번호'): w = 50
        if c.endswith('_출처') or c.endswith('_상태') or c.endswith('_구분') or c == '사람확인_사유': w = 22
        ws.column_dimensions[get_column_letter(j)].width = w
    name_by = {int(v['buildingRosterIdx']): v.get('건물명_확정', '') for v, _ in rows}
    # 근거 시트
    E = wb.create_sheet('근거_주소일치장소')
    eh = ['buildingRosterIdx', '건물명_확정', '검색 도로명주소', '검색 지번주소', '출처', '장소명', '분류', '전화', '링크']
    for j, h in enumerate(eh, 1): c = E.cell(1, j, h); c.font = Font(name=F, bold=True); c.fill = fn
    er = 2
    for i, x in R.items():
        ents = [('카카오맵', n, c, t, u) for n, c, t, u in x.get('K') or []]
        ents += [('114On ' + s(t.get('via')), t['nm'], s(t.get('up')), t['tel'], t['link']) for t in x.get('T') or []]
        for sv, n, c, t, u in ents:
            for j, val in enumerate([i, name_by.get(i, ''), x['road'], x['jibun'], sv, n, c, t, u], 1): E.cell(er, j, val).font = Font(name=F)
            if u: E.cell(er, 6).hyperlink = u; E.cell(er, 6).font = lf; E.cell(er, 9).hyperlink = u; E.cell(er, 9).font = lf
            er += 1
    E.freeze_panes = 'A2'; E.auto_filter.ref = f'A1:I{max(er - 1, 1)}'
    for j, w in enumerate([14, 24, 36, 32, 16, 30, 30, 14, 36], 1): E.column_dimensions[get_column_letter(j)].width = w
    G = wb.create_sheet('근거_KB부동산·집품')
    gh = ['buildingRosterIdx', '건물명_확정', '검색 도로명주소', 'KB 단지명', 'KB 관리사무소 전화', 'KB 세대수', 'KB 링크', '집품 건물명', '집품 건물유형', '집품 링크']
    for j, h in enumerate(gh, 1): c = G.cell(1, j, h); c.font = Font(name=F, bold=True); c.fill = fn
    gr = 2
    for i, x in R.items():
        kb, zp = x.get('kb'), x.get('zp')
        if not kb and not zp: continue
        kbl = 'https://kbland.kr/c/' + str(kb['no']) if kb else ''
        zpl = ('https://zippoom.com/search/' + quote((x['road'] or x['jibun']).replace('인천광역시', '인천')) + '?searchWithAddress=false') if zp else ''
        vals = [i, name_by.get(i, ''), x['road'], (kb or {}).get('main_nm') or (kb or {}).get('nm', ''), (kb or {}).get('tel', ''), (kb or {}).get('hh', ''), kbl, (zp or {}).get('nm', ''), (zp or {}).get('type', ''), zpl]
        for j, val in enumerate(vals, 1): G.cell(gr, j, val).font = Font(name=F)
        if kbl: G.cell(gr, 7).hyperlink = kbl; G.cell(gr, 7).font = lf
        if zpl: G.cell(gr, 10).hyperlink = zpl; G.cell(gr, 10).font = lf
        gr += 1
    G.freeze_panes = 'A2'; G.auto_filter.ref = f'A1:J{max(gr - 1, 1)}'
    for j, w in enumerate([14, 24, 36, 26, 16, 10, 30, 24, 12, 40], 1): G.column_dimensions[get_column_letter(j)].width = w
    H = wb.create_sheet('검수_후보번호')
    hh = ['buildingRosterIdx', '건물명_확정', '도로명주소', '후보 번호', '후보 장소명', '거리 m', '사유', '링크']
    for j, h in enumerate(hh, 1): c = H.cell(1, j, h); c.font = Font(name=F, bold=True); c.fill = fp
    hr = 2
    for i, x in R.items():
        for h in x.get('held') or []:
            for j, val in enumerate([i, name_by.get(i, ''), x['road'], h['v'], h['name'], h.get('d'), h['reason'], h['link']], 1): H.cell(hr, j, val).font = Font(name=F)
            H.cell(hr, 8).hyperlink = h['link']; H.cell(hr, 8).font = lf; hr += 1
    for j, w in enumerate([14, 24, 36, 16, 30, 8, 44, 36], 1): H.column_dimensions[get_column_letter(j)].width = w
    # 지표 시트 · 수식
    S = wb.create_sheet('지표', 0); last = len(rows) + 1
    colL = {c: get_column_letter(j) for j, c in enumerate(cols, 1)}
    rg = lambda c: f"'{sheet}'!${colL[c]}$2:${colL[c]}${last}"
    bd = Border(*(Side(style='thin', color='BFBFBF'),) * 4); sf = PatternFill('solid', fgColor='1F3864')
    S.column_dimensions['A'].width = 52; S.column_dimensions['B'].width = 12; S.column_dimensions['C'].width = 10
    S.cell(1, 1, f'인천 {gu} {dong} 재검색 지표 · {date}').font = Font(name=F, bold=True, size=14)
    S.cell(2, 1, f'{dong} 전체 {len(rows):,}건. 수치는 {sheet} 시트에서 수식으로 계산됨').font = Font(name=F, italic=True, color='595959')
    st = {'r': 4}
    def sec(t):
        for j, h in enumerate([t, '건수', '비율'], 1): c = S.cell(st['r'], j, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = sf; c.border = bd
        st['r'] += 1
    def line(label, f, base=None, bold=False, ind=0):
        r = st['r']
        for j, val in enumerate(['   ' * ind + label, f, f'=IF({base}=0,0,B{r}/{base})' if base else None], 1):
            c = S.cell(r, j, val); c.font = Font(name=F, bold=bold); c.border = bd
        S.cell(r, 2).number_format = '#,##0'; S.cell(r, 3).number_format = '0.0%'
        st['r'] += 1; return r
    sec('처리 범위')
    tot = line(f'{dong} 전체 건물', f'=COUNTA({rg("buildingRosterIdx")})', bold=True)
    for g in 'SABCD': line(f'{g}등급', f'=COUNTIF({rg("등급")},"{g}")', f'$B${tot}', ind=1)
    st['r'] += 1; sec('건물명')
    nt = line('건물명 작업 대상 · (건물명없음)', f'=COUNTIF({rg("건물명")},"(건물명없음)*")', bold=True)
    line('건물명 채움', f'=COUNTIF({rg("건물명_상태")},"채움*")', f'$B${nt}', bold=True)
    for lab, key in [('카카오 도로명주소 건물명', '채움 · 카카오 주소검색'), ('카카오맵 건물 장소', '채움 · 카카오맵 장소'), ('집품', '채움 · 집품'), ('KB부동산 단지명', '채움 · KB부동산')]:
        line(lab, f'=COUNTIF({rg("건물명_상태")},"{key}")', f'$B${nt}', ind=1)
    line('미해결', f'=COUNTIF({rg("건물명_상태")},"미해결")', f'$B${nt}')
    st['r'] += 1; sec('대표번호')
    pt = line('번호 작업 대상 · 대표번호 빈 칸', f'=COUNTIF({rg("번호_상태")},"<>대상아님*")', bold=True)
    line('대표번호 채움', f'=COUNTIF({rg("번호_상태")},"채움*")', f'$B${pt}', bold=True)
    for lab, key in [('카카오맵', '채움 · 카카오맵*'), ('KB부동산 관리사무소', '채움 · KB부동산*'), ('114On', '채움 · 114On*')]:
        line(lab, f'=COUNTIF({rg("번호_상태")},"{key}")', f'$B${pt}', ind=1)
    line('대리 연락처만 확보', f'=COUNTIF({rg("번호_상태")},"대리 연락처만 확보")', f'$B${pt}', bold=True)
    line('미해결 · 대리 연락처도 없음', f'=COUNTIF({rg("번호_상태")},"미해결")', f'$B${pt}')
    st['r'] += 1; sec('검수')
    line('사람 확인 필요', f'=COUNTIF({rg("사람확인필요")},"예")', f'$B${tot}', bold=True)
    line('검수 후보 번호가 있는 건물', f'=COUNTIF({rg("검수_후보번호")},"?*")', f'$B${tot}', ind=1)
    line('시설 이름만 찾은 건물', f'=COUNTIF({rg("시설건물")},"예")', f'$B${tot}', ind=1)
    line('카카오맵 주소 일치 장소가 없는 건물', f'=COUNTIF({rg("주소일치_카카오장소수")},0)', f'$B${tot}')
    st['r'] += 1
    notes = ['참고 사항',
             '· 판정 규칙은 프로젝트 문서 claude/인수인계.md 와 claude/pipeline/common.py 를 따름',
             '· 대표번호는 건물 자체 번호나 관리사무소 번호만. 건물 안 점포와 주변 공인중개사는 대리 연락처 칸에만 적음',
             '· 검수_후보번호는 자동으로 넣지 않은 후보. 이름 불일치, 네이버 위치 불일치나 미등록, 다른 시설 관리사무소 등',
             '· 파란 밑줄 셀을 누르면 해당 출처 페이지가 열림']
    for i, t in enumerate(notes): S.cell(st['r'], 1, t).font = Font(name=F, bold=(i == 0), color='404040'); st['r'] += 1
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, f'콜리스트_{gu}_{dong}_재검색결과_{date}.xlsx')
    wb.save(path)
    if shutil.which('soffice'):
        tmp = os.path.join(out, '_recalc'); os.makedirs(tmp, exist_ok=True)
        subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path], capture_output=True, timeout=180)
        rc = os.path.join(tmp, os.path.basename(path))
        if os.path.exists(rc): shutil.move(rc, path)
        shutil.rmtree(tmp, ignore_errors=True)
    return path, len(rows)

# ---------------- HTML ----------------
CSS = """
:root{--bg:#f6f7f9;--card:#fff;--ink:#1b1f24;--sub:#5b6470;--line:#e3e6ea;--acc:#1c6b86;--accbg:#e1eff4;--warn:#9a5b0b;--warnbg:#fbeedb;--ok:#2e7d50;--okbg:#e3f2e8;--chip:#f1f3f5;--sec:#1b2530}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){color-scheme:dark;--bg:#111418;--card:#1a1e24;--ink:#e8eaed;--sub:#9aa3ad;--line:#2b3139;--acc:#6cc0dc;--accbg:#16323d;--warn:#ffb86b;--warnbg:#3a2a16;--ok:#74c795;--okbg:#173226;--chip:#242a31;--sec:#e8eaed}}
:root[data-theme="dark"]{color-scheme:dark;--bg:#111418;--card:#1a1e24;--ink:#e8eaed;--sub:#9aa3ad;--line:#2b3139;--acc:#6cc0dc;--accbg:#16323d;--warn:#ffb86b;--warnbg:#3a2a16;--ok:#74c795;--okbg:#173226;--chip:#242a31;--sec:#e8eaed}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:"Pretendard","Malgun Gothic","Apple SD Gothic Neo",sans-serif;font-size:14px;line-height:1.5}
header{padding:24px 16px 8px;max-width:1280px;margin:0 auto}h1{font-size:20px;margin:0 0 6px}.desc{color:var(--sub);margin:0 0 10px;max-width:90ch}
.sum{display:flex;gap:8px;flex-wrap:wrap;margin:6px 0 4px}.sum span{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:6px 10px;font-size:13px}
.bar{position:sticky;top:0;background:var(--bg);padding:10px 16px;z-index:5;border-bottom:1px solid var(--line)}
.bar-in{max-width:1280px;margin:0 auto;display:flex;gap:8px;align-items:center;flex-wrap:wrap}
input[type=search]{flex:1;min-width:200px;padding:9px 12px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--ink);font-size:14px}
select{padding:8px 10px;border:1px solid var(--line);border-radius:8px;background:var(--card);color:var(--ink);font-size:13px}
label.f{font-size:13px;color:var(--sub);display:flex;gap:4px;align-items:center}
.count{color:var(--sub);font-size:13px}
main{max-width:1280px;margin:0 auto;padding:12px 16px 48px}
h2.sec{font-size:17px;margin:22px 0 8px;color:var(--sec)}h3.gr{font-size:14px;margin:14px 0 6px;color:var(--sub)}
.row{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px;margin-bottom:10px;display:grid;grid-template-columns:44px minmax(0,1fr) auto;gap:12px;align-items:start}
.row.h{border-left:4px solid var(--warn)}
.no{color:var(--sub);font-variant-numeric:tabular-nums;padding-top:2px}
.ttl{font-weight:700;font-size:15px}.addr{font-weight:600}.jib{color:var(--sub);font-size:13px}
.f1{margin-top:6px;font-size:13px;display:grid;grid-template-columns:88px minmax(0,1fr);gap:2px 10px}.f1 b{color:var(--sub);font-weight:600}
.f1 a,.addr a,.jib a{color:var(--acc)}
.meta{display:flex;gap:6px;flex-wrap:wrap;margin-top:8px}
.chip{background:var(--chip);border-radius:999px;padding:2px 9px;font-size:12px;color:var(--sub)}
.chip.warn{background:var(--warnbg);color:var(--warn)}.chip.ok{background:var(--okbg);color:var(--ok)}.chip.new{background:var(--accbg);color:var(--acc)}
.px{margin-top:6px;font-size:12.5px;color:var(--sub)}
.links{display:flex;gap:6px;flex-wrap:wrap;justify-content:flex-end;max-width:420px}
.links a,.links button{display:inline-block;padding:6px 10px;border-radius:7px;background:var(--accbg);color:var(--acc);text-decoration:none;font-size:12.5px;font-weight:600;border:0;cursor:pointer;font-family:inherit}
.links a:focus-visible,.links button:focus-visible{outline:2px solid var(--acc);outline-offset:2px}
.links .off{opacity:.45;pointer-events:none}
@media (max-width:760px){.row{grid-template-columns:minmax(0,1fr)}.no{display:none}.links{justify-content:flex-start;max-width:none}.f1{grid-template-columns:minmax(0,1fr)}}
"""

JS = """
const q=document.getElementById('q'),rows=[...document.querySelectorAll('.row')],cnt=document.getElementById('cnt'),gsel=document.getElementById('g'),fsel=document.getElementById('flt');
function f(){const t=q.value.trim().toLowerCase(),g=gsel?gsel.value:'',k=fsel?fsel.value:'';let n=0;
rows.forEach(r=>{const ok=(!t||r.dataset.s.includes(t))&&(!g||r.dataset.g===g)&&(!k||r.dataset.k.split(' ').includes(k));r.hidden=!ok;if(ok)n++;});
document.querySelectorAll('[data-grp]').forEach(h=>{const sel=h.dataset.grp;const exact=h.tagName==='H3';h.hidden=!rows.some(r=>!r.hidden&&(exact?r.dataset.in===sel:r.dataset.in.startsWith(sel)));});
cnt.textContent=n+' / '+rows.length+'건';}
[q,gsel,fsel].forEach(e=>e&&e.addEventListener('input',f));f();
document.querySelectorAll('button[data-copy]').forEach(b=>b.addEventListener('click',()=>{const v=b.dataset.copy;const done=()=>{const o=b.textContent;b.textContent='복사됨';setTimeout(()=>b.textContent=o,1200);};
const fb=()=>{const i=document.createElement('input');i.value=v;document.body.appendChild(i);i.select();try{document.execCommand('copy');done();}catch(e){}i.remove();};
if(navigator.clipboard){navigator.clipboard.writeText(v).then(done).catch(fb);}else fb();}));
"""

E = html.escape

def a(url, text):
    return f'<a href="{E(url)}" target="_blank" rel="noopener">{E(text)}</a>' if url else E(text)

def card(x, n, mode):
    L = links_for(x); q = x['road'] or x['jibun']; flags = x.get('flags') or []
    chips = [f'<span class="chip">{E(x["grade"] or "등급없음")}{"등급" if x["grade"] else ""}</span>', f'<span class="chip">idx {x["i"]}</span>']
    if s(x['hh']) or s(x['ho']): chips.append(f'<span class="chip">세대 {E(s(x["hh"]) or "0")} · 호 {E(s(x["ho"]) or "0")}</span>')
    kinds = []
    if x['nt'] and x.get('nm'): chips.append('<span class="chip new">건물명 채움</span>'); kinds.append('name')
    if x['pt'] and x.get('ph'): chips.append('<span class="chip ok">대표번호 채움</span>'); kinds.append('phone')
    elif x['pt'] and (x.get('st') or x.get('re')): chips.append('<span class="chip">대리 연락처만</span>'); kinds.append('proxy')
    for fl in flags: chips.append(f'<span class="chip warn">{E(fl)}</span>')
    if flags: kinds.append('check')
    if x.get('fac') and x['fac'][0]: chips.append(f'<span class="chip warn">시설 건물 · {E(x["fac"][0])}</span>')
    chips.append(f'<span class="chip">카카오 주소일치 장소 {x.get("km", 0)}곳</span>')
    if x.get('tm'): chips.append(f'<span class="chip">114On 점포 {x["tm"]}곳</span>')
    # 제목 줄 · 입력한 정보 / 건물명
    nm_txt = (x['nm']['v'] + ' · 새로 채움') if (x['nt'] and x.get('nm')) else ((x['n0'] + ' · 기존 건물명') if x['n0'] else '건물명 없음')
    if x['pt'] and x.get('ph'): lead = '대표번호 ' + x['ph']['v']
    elif x['pt'] and (x.get('st') or x.get('re')): lead = f"대리 연락처 {len(x.get('st') or []) + len(x.get('re') or [])}곳"
    elif x['nt'] and x.get('nm'): lead = '건물명'
    else: lead = '채운 정보 없음'
    body = []
    if mode == 'filled':
        body.append(f'<div class="ttl">{E(lead)} / {E(nm_txt)}</div>')
    body.append(f'<div class="addr">{E(x["road"] or x["jibun"])} <span class="jib">{a(L["카카오맵"], "카카오맵")} · {a(L["네이버지도"], "네이버지도")}</span></div>')
    if x['road'] and x['jibun']: body.append(f'<div class="jib">지번 {E(short_addr(x["jibun"]))}</div>')
    if mode != 'filled' and not x['nt'] and x['n0']: body.append(f'<div class="jib">건물명 {E(x["n0"])}</div>')
    f1 = []
    if x['nt']:
        f1.append(('건물명', a(x['nm']['link'], x['nm']['v']) + f' <span class="jib">{E(x["nm"]["src"])}</span>' if x.get('nm') else '미해결'))
    elif mode == 'filled':
        f1.append(('건물명', '변경 없음'))
    if x['pt']:
        if x.get('ph'): f1.append(('대표번호', a(x['ph']['link'], x['ph']['v']) + f' <span class="jib">{E(x["ph"]["name"])} · {E(x["ph"]["src"])}</span>'))
        else: f1.append(('대표번호', '없음'))
        pr = proxy_rows(x)
        if pr: f1.append(('대리 연락처', '<br>'.join(a(link, f'{nm} {p}') + f' <span class="jib">{E(g)}</span>' for g, nm, p, src, link in pr)))
        elif not x.get('ph'): f1.append(('대리 연락처', '없음'))
    for h in x.get('held') or []:
        f1.append(('검수 후보', a(h['link'], f"{h['v']} {h['name']}") + f' <span class="jib">{E(h["reason"])}' + (f' · {h["d"]}m' if h.get('d') is not None else '') + '</span>'))
    if f1: body.append('<div class="f1">' + ''.join(f'<b>{E(k)}</b><div>{v}</div>' for k, v in f1) + '</div>')
    body.append(f'<div class="meta">{"".join(chips)}</div>')
    ls = [f'<a href="{E(L[k])}" target="_blank" rel="noopener">{k}</a>' if L[k] else f'<a class="off">{k}</a>' for k in ['카카오 로드뷰', '카카오맵', '네이버지도', 'KB부동산', '집품', '114On']]
    ls.append(f'<button type="button" data-copy="{E(q)}">주소 복사</button>')
    srch = ' '.join([x['road'], x['jibun'], (x.get('nm') or {}).get('v', ''), x['n0'], str(x['i']), x['grade']] + flags).lower()
    sec = 'h' if x.get('human') else 'r'
    return (f'<div class="row{" h" if x.get("human") else ""}" data-s="{E(srch)}" data-g="{E(x["grade"])}" data-k="{" ".join(kinds)}" data-in="{sec}{E(x["grade"] or "-")}">'
            f'<div class="no">{n}</div><div>{"".join(body)}</div><div class="links">{"".join(ls)}</div></div>')

def page(items, title, desc, mode, path):
    items = sorted(items, key=sort_key)
    human = [x for x in items if x.get('human')]; rest = [x for x in items if not x.get('human')]
    parts, n = [], 0
    for label, sec, arr in (('사람 확인 필요', 'h', human), ('나머지', 'r', rest)):
        parts.append(f'<h2 class="sec" data-grp="{sec}" >{label} · {len(arr)}건</h2>' if arr else '')
        grades = sorted({x['grade'] for x in arr}, key=lambda g: GRADE_ORDER.get(s(g).upper(), 9))
        for g in grades:
            sub = [x for x in arr if x['grade'] == g]
            parts.append(f'<h3 class="gr" data-grp="{sec}{E(g or "-")}">{E(g or "등급 없음")}{"등급" if g else ""} · {len(sub)}건</h3>')
            for x in sub:
                n += 1; parts.append(card(x, n, mode))
    gsel = '<select id="g" aria-label="등급"><option value="">전체 등급</option>' + ''.join(f'<option value="{g}">{g}등급</option>' for g in 'SABCD') + '</select>'
    fsel = ('<select id="flt" aria-label="채운 항목"><option value="">전체</option><option value="name">건물명 채움</option><option value="phone">대표번호 채움</option><option value="proxy">대리 연락처만</option><option value="check">검수 필요만</option></select>') if mode == 'filled' else ''
    summ = f'<div class="sum"><span>전체 {len(items)}건</span><span>사람 확인 필요 {len(human)}건</span><span>나머지 {len(rest)}건</span></div>'
    doc = f"""<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{E(title)}</title><style>{CSS}</style></head><body>
<header><h1>{E(title)}</h1><p class="desc">{E(desc)}</p>{summ}</header>
<div class="bar"><div class="bar-in"><input id="q" type="search" placeholder="주소, 지번, 건물명, idx, 검수 사유로 찾기" aria-label="검색">{gsel}{fsel}<span class="count" id="cnt"></span></div></div>
<main>{''.join(parts) if items else '<p class="desc">해당하는 건물이 없습니다.</p>'}</main><script>{JS}</script></body></html>"""
    with open(path, 'w', encoding='utf-8') as f: f.write(doc)
    return len(items), len(human), len(rest)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', required=True); ap.add_argument('--work', default='work')
    ap.add_argument('--gu', required=True); ap.add_argument('--dong', required=True)
    ap.add_argument('--out', default='out'); ap.add_argument('--date', required=True)
    a_ = ap.parse_args()
    d = os.path.join(a_.work, f'{a_.gu}_{a_.dong}')
    res = load(os.path.join(d, 'results.json'), [])
    if not res: raise SystemExit('results.json 이 없음. merge.py 먼저')
    R = {int(x['i']): x for x in res}
    df = None
    for enc in ('utf-8-sig', 'cp949'):
        try: df = pd.read_csv(a_.src, encoding=enc); break
        except UnicodeDecodeError: continue
    df = df[(df['시군구'] == a_.gu) & (df['법정동'] == a_.dong)].reset_index(drop=True)
    xp, nx = build_xlsx(df, R, a_.gu, a_.dong, a_.out, a_.date)
    print('엑셀', xp, nx, '행')
    nmU = lambda x: x['nt'] and not x.get('nm')
    phU = lambda x: x['pt'] and not x.get('ph')
    sets = {
        '사람확인필요': [x for x in res if x.get('human')],
        '건물명번호_둘다미해결': [x for x in res if nmU(x) and phU(x)],
        '건물명만미해결': [x for x in res if nmU(x) and not phU(x)],
        '번호만미해결': [x for x in res if phU(x) and not nmU(x)],
        '채운정보_전체': [x for x in res if (x['nt'] and x.get('nm')) or (x['pt'] and (x.get('ph') or x.get('st') or x.get('re')))],
    }
    base = f'인천 {a_.gu} {a_.dong} 재검색 결과 · {a_.date}. 사람 확인이 필요한 주소를 맨 위에 두고, 두 구역 모두 등급 S, A, B, C, D 순으로 정렬했습니다. 버튼은 도로명주소로 검색하고, 도로명이 없으면 지번으로 검색합니다.'
    descs = {
        '사람확인필요': base + ' 건물명 미해결, 대리 연락처도 없는 건물, 건물명 검수 필요, 검수 후보 번호가 있는 건물처럼 로드뷰나 지도를 직접 봐야 하는 주소입니다.',
        '건물명번호_둘다미해결': base + ' 건물명도 못 찾았고 대표번호도 없는 건물입니다. 대리 연락처가 있어도 대표번호가 아니면 번호 미해결로 봤습니다.',
        '건물명만미해결': base + ' 건물명은 못 찾았지만 대표번호는 있는 건물입니다.',
        '번호만미해결': base + ' 건물명은 있지만 대표번호가 없는 건물입니다.',
        '채운정보_전체': base + ' 건물명, 대표번호, 대리 연락처 중 무엇이든 새로 채운 주소 전부입니다. 제목 줄은 입력한 정보 / 건물명 순서이고, 값마다 출처 링크가 달려 있습니다.',
    }
    titles = {'사람확인필요': '사람 확인 필요 목록', '건물명번호_둘다미해결': '건물명·대표번호 둘 다 미해결 목록', '건물명만미해결': '건물명만 미해결 목록', '번호만미해결': '대표번호만 미해결 목록', '채운정보_전체': '채운 정보 전체 목록'}
    for k, items in sets.items():
        path = os.path.join(a_.out, f'{a_.gu}_{a_.dong}_{k}_목록_{a_.date}.html')
        n, h, r_ = page(items, f'{a_.dong} {titles[k]}', descs[k], 'filled' if k == '채운정보_전체' else 'list', path)
        assert h + r_ == n
        print(f'HTML {k}: {n}건 · 사람 확인 필요 {h} · 나머지 {r_} → {path}')
    A, B, C = ({x['i'] for x in sets[k]} for k in ('건물명번호_둘다미해결', '건물명만미해결', '번호만미해결'))
    union_ok = len(A | B | C) == sum(1 for x in res if nmU(x) or phU(x))
    print('검증 · 세 미해결 목록 겹침', len(A & B), len(A & C), len(B & C), '· 합집합이 공백 건물 수와 같음' if union_ok else '· 합집합 불일치')

if __name__ == '__main__':
    main()
