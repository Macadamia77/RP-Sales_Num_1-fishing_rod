"""선택한 행만 결과 파일로 · build_outputs.py 를 고치지 않고 감싸서 쓴다 (이 파일은 원본이 아니라 새로 만든 어댑터)

  · 작업 폴더의 source.csv(이번 실행에서 선택된 조회 대상 행)와 results.json 으로 build_outputs 의 build_xlsx·page 를 그대로 부른다
    → 엑셀 1개(지표·결과·근거 시트)와 HTML 5종: 사람확인필요, 건물명번호_둘다미해결, 건물명만미해결, 번호만미해결, 채운정보_전체
    HTML 제목·설명은 Codex 버전 결과물과 같은 형식(「<구> <동> <목록>」, 「<날짜> · 이번 실행에서 선택한 행만 표시합니다…」)
  · 수식 차단 · 조회로 가져온 이름이 「=」로 시작하면 openpyxl이 엑셀 수식으로 저장한다.
    지표 시트(일부러 쓴 수식)를 뺀 모든 시트에서 수식 칸을 글자로 바꾼다
  · 수식 차단을 먼저 한 뒤에 LibreOffice로 지표 수식을 다시 계산한다 (LibreOffice가 없으면 건너뜀)

사용 예
  python export_selected.py --work <data 폴더> --gu 미추홀구 --dong 관교동 --out <결과 폴더> --date 2026-09-29
"""
import argparse
import os
import shutil
import subprocess
from unittest import mock

import pandas as pd
from openpyxl import load_workbook

import build_outputs

FORMULA_SHEETS = ('지표',)


def guard_formulas(path):
    """지표 시트 밖의 수식 칸을 글자로. 바꾼 칸 수를 돌려줌"""
    wb = load_workbook(path)
    n = 0
    for ws in wb.worksheets:
        if ws.title in FORMULA_SHEETS:
            continue
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == 'f':
                    c.data_type = 's'
                    c.quotePrefix = True
                    n += 1
    if n:
        wb.save(path)
    return n


def recalc(path):
    """LibreOffice로 한 번 열어 수식 값을 채워 저장. 실패해도 원래 파일은 그대로"""
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    if not exe:
        return False
    out = os.path.dirname(path) or '.'
    tmp = os.path.join(out, '_recalc')
    os.makedirs(tmp, exist_ok=True)
    try:
        subprocess.run([exe, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, path], capture_output=True, timeout=180)
        rc = os.path.join(tmp, os.path.basename(path))
        if os.path.exists(rc):
            shutil.move(rc, path)
            return True
        return False
    except (OSError, subprocess.SubprocessError):
        return False
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


SETS = ['사람확인필요', '건물명번호_둘다미해결', '건물명만미해결', '번호만미해결', '채운정보_전체']


def pick(res):
    """build_outputs.main 과 같은 분류"""
    nmU = lambda x: x['nt'] and not x.get('nm')  # noqa: E731
    phU = lambda x: x['pt'] and not x.get('ph')  # noqa: E731
    return {
        '사람확인필요': [x for x in res if x.get('human')],
        '건물명번호_둘다미해결': [x for x in res if nmU(x) and phU(x)],
        '건물명만미해결': [x for x in res if nmU(x) and not phU(x)],
        '번호만미해결': [x for x in res if phU(x) and not nmU(x)],
        '채운정보_전체': [x for x in res if (x['nt'] and x.get('nm')) or (x['pt'] and (x.get('ph') or x.get('st') or x.get('re')))],
    }, nmU, phU


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--work', required=True)
    ap.add_argument('--gu', required=True)
    ap.add_argument('--dong', required=True)
    ap.add_argument('--out', required=True)
    ap.add_argument('--date', required=True)
    a = ap.parse_args()
    d = os.path.join(a.work, f'{a.gu}_{a.dong}')
    src = os.path.join(d, 'source.csv')
    res = build_outputs.load(os.path.join(d, 'results.json'), [])
    if not os.path.exists(src) or not res:
        raise SystemExit(f'source.csv 나 results.json 이 없음 · {d}')
    R = {int(x['i']): x for x in res}
    df = pd.read_csv(src, encoding='utf-8-sig')
    df = df[(df['시군구'] == a.gu) & (df['법정동'] == a.dong)].reset_index(drop=True)
    os.makedirs(a.out, exist_ok=True)
    # build_outputs 안의 LibreOffice 재계산은 수식 차단 뒤로 미룸
    with mock.patch.object(build_outputs.shutil, 'which', return_value=None):
        xlsx, nx = build_outputs.build_xlsx(df, R, a.gu, a.dong, a.out, a.date)
    n = guard_formulas(xlsx)
    ok = recalc(xlsx)
    print(f'엑셀 {xlsx} {nx}행 · 수식 차단 {n}칸 · 지표 재계산 {"완료" if ok else "건너뜀 (LibreOffice 없음 또는 실패)"}')
    sets, nmU, phU = pick(res)
    desc = f'{a.date} · 이번 실행에서 선택한 행만 표시합니다. 대리 연락처는 대표번호와 구분합니다.'
    for k in SETS:
        path = os.path.join(a.out, f'{a.gu}_{a.dong}_{k}_목록_{a.date}.html')
        cnt, h, r_ = build_outputs.page(sets[k], f'{a.gu} {a.dong} {k}', desc, 'filled' if k == '채운정보_전체' else 'list', path)
        assert h + r_ == cnt
        print(f'HTML {k}: {cnt}건 · 사람 확인 필요 {h} · 나머지 {r_}')
    A, B, C = ({x['i'] for x in sets[k]} for k in ('건물명번호_둘다미해결', '건물명만미해결', '번호만미해결'))
    union_ok = len(A | B | C) == sum(1 for x in res if nmU(x) or phU(x))
    if (A & B) or (A & C) or (B & C) or not union_ok:
        raise SystemExit('검증 실패 · 미해결 목록이 겹치거나 합집합이 공백 건물 수와 다름')
    print('검증 · 세 미해결 목록 겹침 없음 · 합집합이 공백 건물 수와 같음')


if __name__ == '__main__':
    main()
