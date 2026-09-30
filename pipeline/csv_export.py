"""엑셀 결과 시트 → CSV · 하이퍼링크는 바로 옆 「<열>_링크」 열로 보존"""
import csv
import os

from openpyxl import load_workbook

LINK_COLS = ['건물명_확정', '대표번호_확정'] + [f'대리{n}_{k}' for n in range(1, 6) for k in ('건물명', '번호')]


def xlsx_to_csv(xlsx, dong):
    wb = load_workbook(xlsx)
    sheet = f'{dong}_결과'[:31]
    if sheet not in wb.sheetnames:
        raise ValueError(f'{os.path.basename(xlsx)} 에 {sheet} 시트가 없음')
    ws = wb[sheet]
    head = [c.value for c in ws[1]]
    cols = []
    for j, h in enumerate(head):
        cols.append((j, h, False))
        if h in LINK_COLS:
            cols.append((j, f'{h}_링크', True))
    path = xlsx[:-5] + '.csv'
    tmp = path + '.tmp'
    n = 0
    with open(tmp, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow([c[1] for c in cols])
        for row in ws.iter_rows(min_row=2):
            out = []
            for j, _, is_link in cols:
                c = row[j]
                if is_link:
                    out.append(c.hyperlink.target if c.hyperlink and c.hyperlink.target else '')
                else:
                    out.append('' if c.value is None else c.value)
            w.writerow(out)
            n += 1
    os.replace(tmp, path)
    return path, n
